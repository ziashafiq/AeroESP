from collections import Counter, defaultdict
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

from research_review.models import (
    ResearchExperiment,
    ResearchQuestion,
    ResearchRun,
)


EXPECTED_ROUNDS = ["R003", "R004", "R005", "R006", "R007"]

EXPECTED_SKILLS = {
    1: "main_idea",
    2: "specific_information",
    3: "specific_information",
    4: "inference",
    5: "inference",
    6: "vocabulary_in_context",
    7: "technical_comprehension",
    8: "technical_comprehension",
    9: "cause_and_effect",
    10: "applied_understanding",
}

# Frozen V4 run-slot design.
EXPECTED_RUNS = {
    1: {
        "ai_model": "ChatGPT",
        "service": "ChatGPT",
        "condition": "Free Simple / Think OFF",
        "displayed_model": "NOT DISCLOSED",
    },
    2: {
        "ai_model": "DeepSeek",
        "service": "DeepSeek",
        "condition": "Free Instant / DeepThink OFF / Search OFF",
        "displayed_model": "NOT DISCLOSED",
    },
    3: {
        "ai_model": "Grok",
        "service": "Grok",
        "condition": "Free Fast",
        "displayed_model": "NOT DISCLOSED",
    },
    4: {
        "ai_model": "Claude",
        "service": "Claude",
        "condition": "Sonnet 5 / Medium",
        "displayed_model": "Sonnet 5",
    },
    5: {
        "ai_model": "Gemini",
        "service": "Gemini",
        "condition": "3.6 Flash / Extended Thinking OFF",
        "displayed_model": "3.6 Flash",
    },
    6: {
        "ai_model": "Qwen",
        "service": "Qwen",
        "condition": "3.7 Plus / Auto",
        "displayed_model": "3.7 Plus",
    },
    7: {
        "ai_model": "Copilot",
        "service": "Microsoft Copilot",
        "condition": "Smart",
        "displayed_model": "NOT DISCLOSED",
    },
    8: {
        "ai_model": "Mistral",
        "service": "Mistral",
        "condition": "Vibe Chat / Fast",
        "displayed_model": "Vibe Chat",
    },
}

EXPECTED_QUESTION_IDS = {
    f"AEV4-{i:04d}"
    for i in range(1, 401)
}

EXPECTED_RECOVERY_COUNTS = {
    "ORIGINAL_COMBINED": 340,
    "RECOVERED_ORIGINAL_CHAT": 60,
}

REQUIRED_COLUMNS = [
    # Core item data
    "R_Number",
    "AI_Model",
    "Domain",
    "Topic",
    "Protocol",
    "Experiment_ID",
    "CEFR",
    "Cognitive_Level",
    "Item_Number",
    "Skill",
    "Stem",
    "Option_A",
    "Option_B",
    "Option_C",
    "Option_D",
    "Correct_Answer",
    "Explanation",
    "Source_Evidence",

    # Canonical provenance / recovery metadata
    "Raw_Format",
    "JSON_Valid",
    "Schema_Compliant",
    "Recovery_Status",
    "Recovery_Source",

    # Frozen blind-ID map metadata
    "QUESTION_ID",
    "INTERNAL_SEQ",
    "RUN_SLOT",
    "SERVICE",
    "CONDITION",
    "ID_MAP_RAW_OUTPUT_FORMAT",
]


class Command(BaseCommand):
    help = (
        "Audit and import the frozen AeroESP V4 Main Study "
        "(R003-R007 only) into research_review using the "
        "authoritative pre-existing blinded QUESTION_ID values."
    )

    def add_arguments(self, parser):
        parser.add_argument(
            "xlsx_path",
            type=str,
            help=(
                "Path to "
                "AeroESP_V4_CANONICAL_400_WITH_IDS.xlsx"
            ),
        )
        parser.add_argument(
            "--apply",
            action="store_true",
            help=(
                "Actually write data. Without this flag, "
                "perform audit only."
            ),
        )

    def handle(self, *args, **options):
        if load_workbook is None:
            raise CommandError(
                "openpyxl is not installed. Run: "
                "python -m pip install openpyxl"
            )

        path = Path(
            options["xlsx_path"]
        ).expanduser().resolve()

        if not path.exists():
            raise CommandError(
                f"File not found: {path}"
            )

        self.stdout.write(
            f"Reading: {path}"
        )

        rows = self._read_excel(path)

        self.stdout.write(
            f"Rows found in workbook: {len(rows)}"
        )

        all_round_counts = Counter(
            row["R_Number"]
            for row in rows
        )

        self.stdout.write(
            "\nWORKBOOK ROUND COUNTS"
        )
        for round_id in sorted(
            all_round_counts
        ):
            self.stdout.write(
                f"  {round_id}: "
                f"{all_round_counts[round_id]}"
            )

        unexpected_rounds = sorted(
            set(all_round_counts)
            - set(EXPECTED_ROUNDS)
        )
        if unexpected_rounds:
            raise CommandError(
                "This importer accepts only the "
                "frozen V4 Main Study R003-R007. "
                f"Unexpected rounds: "
                f"{unexpected_rounds}"
            )

        self._audit(rows)

        self.stdout.write(
            self.style.SUCCESS(
                "\nAUDIT PASSED - "
                "V4 Canonical 400 + frozen Blind IDs "
                "are structurally valid."
            )
        )

        if not options["apply"]:
            self.stdout.write(
                self.style.WARNING(
                    "\nDRY RUN ONLY - "
                    "database was NOT modified."
                )
            )
            self.stdout.write(
                "Run again with --apply only "
                "after reviewing this audit."
            )
            return

        self._import(rows)

    @staticmethod
    def _clean(value):
        if isinstance(value, str):
            return value.strip()
        return value

    @staticmethod
    def _to_int(value, field_name):
        try:
            return int(value)
        except (TypeError, ValueError):
            raise CommandError(
                f"Invalid {field_name}: {value!r}"
            )

    @staticmethod
    def _to_bool(value, field_name):
        if isinstance(value, bool):
            return value

        if isinstance(value, int):
            if value in (0, 1):
                return bool(value)

        if isinstance(value, str):
            normalized = value.strip().lower()

            if normalized in {
                "true", "1", "yes", "y", "pass"
            }:
                return True

            if normalized in {
                "false", "0", "no", "n", "fail"
            }:
                return False

        raise CommandError(
            f"Invalid boolean value for "
            f"{field_name}: {value!r}"
        )

    def _read_excel(self, path):
        workbook = load_workbook(
            path,
            read_only=True,
            data_only=True,
        )

        if "CANONICAL_400" in workbook.sheetnames:
            sheet = workbook["CANONICAL_400"]
        else:
            sheet = workbook.active

        raw_headers = [
            cell.value
            for cell in next(
                sheet.iter_rows(
                    min_row=1,
                    max_row=1,
                )
            )
        ]

        missing = [
            column
            for column in REQUIRED_COLUMNS
            if column not in raw_headers
        ]

        if missing:
            raise CommandError(
                "This is not the final "
                "WITH_IDS canonical workbook. "
                f"Missing required columns: "
                f"{missing}"
            )

        header_map = {
            name: raw_headers.index(name)
            for name in REQUIRED_COLUMNS
        }

        rows = []

        for excel_row in sheet.iter_rows(
            min_row=2,
            values_only=True,
        ):
            if not any(
                value is not None
                for value in excel_row
            ):
                continue

            row = {}

            for column in REQUIRED_COLUMNS:
                value = excel_row[
                    header_map[column]
                ]
                row[column] = self._clean(
                    value
                )

            for field in [
                "R_Number",
                "AI_Model",
                "Domain",
                "Topic",
                "Protocol",
                "Experiment_ID",
                "CEFR",
                "Cognitive_Level",
                "Skill",
                "Correct_Answer",
                "Raw_Format",
                "Recovery_Status",
                "Recovery_Source",
                "QUESTION_ID",
                "SERVICE",
                "CONDITION",
                "ID_MAP_RAW_OUTPUT_FORMAT",
            ]:
                if row[field] is not None:
                    row[field] = str(
                        row[field]
                    ).strip()

            if (
                row["Cognitive_Level"]
                is not None
            ):
                row["Cognitive_Level"] = (
                    row[
                        "Cognitive_Level"
                    ].upper()
                )

            if (
                row["Correct_Answer"]
                is not None
            ):
                row["Correct_Answer"] = (
                    row[
                        "Correct_Answer"
                    ].upper()
                )

            row["Item_Number"] = (
                self._to_int(
                    row["Item_Number"],
                    "Item_Number",
                )
            )
            row["INTERNAL_SEQ"] = (
                self._to_int(
                    row["INTERNAL_SEQ"],
                    "INTERNAL_SEQ",
                )
            )
            row["RUN_SLOT"] = (
                self._to_int(
                    row["RUN_SLOT"],
                    "RUN_SLOT",
                )
            )

            row["JSON_Valid"] = (
                self._to_bool(
                    row["JSON_Valid"],
                    "JSON_Valid",
                )
            )
            row["Schema_Compliant"] = (
                self._to_bool(
                    row["Schema_Compliant"],
                    "Schema_Compliant",
                )
            )

            rows.append(row)

        return rows

    def _audit(self, rows):
        if len(rows) != 400:
            raise CommandError(
                "Expected exactly 400 "
                "R003-R007 rows; "
                f"found {len(rows)}."
            )

        # --------------------------------------------------
        # Global frozen ID integrity
        # --------------------------------------------------
        question_ids = [
            row["QUESTION_ID"]
            for row in rows
        ]

        if len(set(question_ids)) != 400:
            duplicates = [
                qid
                for qid, count
                in Counter(
                    question_ids
                ).items()
                if count > 1
            ]
            raise CommandError(
                "QUESTION_ID values must "
                "be unique. Duplicates: "
                f"{duplicates[:20]}"
            )

        if set(question_ids) != (
            EXPECTED_QUESTION_IDS
        ):
            missing_ids = sorted(
                EXPECTED_QUESTION_IDS
                - set(question_ids)
            )
            extra_ids = sorted(
                set(question_ids)
                - EXPECTED_QUESTION_IDS
            )
            raise CommandError(
                "Frozen QUESTION_ID set "
                "does not match exactly "
                "AEV4-0001..AEV4-0400.\n"
                f"Missing: "
                f"{missing_ids[:20]}\n"
                f"Extra: "
                f"{extra_ids[:20]}"
            )

        internal_seq = [
            row["INTERNAL_SEQ"]
            for row in rows
        ]

        if set(internal_seq) != (
            set(range(1, 401))
        ):
            raise CommandError(
                "INTERNAL_SEQ must be "
                "a unique permutation "
                "of 1..400."
            )

        if len(set(internal_seq)) != 400:
            raise CommandError(
                "INTERNAL_SEQ contains "
                "duplicates."
            )

        self.stdout.write(
            "\nBLIND-ID QC"
        )
        self.stdout.write(
            "  QUESTION_ID: "
            "400 unique / exact "
            "AEV4-0001..AEV4-0400 -> PASS"
        )
        self.stdout.write(
            "  INTERNAL_SEQ: "
            "1..400 unique -> PASS"
        )

        # --------------------------------------------------
        # Main-study round / protocol integrity
        # --------------------------------------------------
        round_counts = Counter(
            row["R_Number"]
            for row in rows
        )

        self.stdout.write(
            "\nMAIN-STUDY ROUND COUNTS"
        )

        for round_id in EXPECTED_ROUNDS:
            count = round_counts[
                round_id
            ]

            self.stdout.write(
                f"  {round_id}: {count}"
            )

            if count != 80:
                raise CommandError(
                    f"{round_id} must "
                    "contain 80 questions; "
                    f"found {count}."
                )

        protocols = {
            row["Protocol"]
            for row in rows
        }
        if protocols != {
            "AEROESP_GEN_V4"
        }:
            raise CommandError(
                "Unexpected protocol "
                f"values: {protocols}"
            )

        cefr_values = {
            row["CEFR"]
            for row in rows
        }
        if cefr_values != {"B2"}:
            raise CommandError(
                "Unexpected CEFR "
                f"values: {cefr_values}"
            )

        cognitive_values = {
            row["Cognitive_Level"]
            for row in rows
        }
        if cognitive_values != {
            "UNDERSTAND"
        }:
            raise CommandError(
                "Unexpected cognitive "
                f"levels: "
                f"{cognitive_values}"
            )

        # --------------------------------------------------
        # Item / run integrity
        # --------------------------------------------------
        grouped = defaultdict(list)

        for row in rows:
            required_text_fields = [
                "R_Number",
                "AI_Model",
                "Domain",
                "Topic",
                "Protocol",
                "Experiment_ID",
                "Skill",
                "Stem",
                "Option_A",
                "Option_B",
                "Option_C",
                "Option_D",
                "Correct_Answer",
                "Raw_Format",
                "Recovery_Status",
                "QUESTION_ID",
                "SERVICE",
                "CONDITION",
            ]

            for field in (
                required_text_fields
            ):
                value = row[field]

                if (
                    value is None
                    or str(value).strip()
                    == ""
                ):
                    raise CommandError(
                        "Blank required "
                        f"field '{field}' "
                        f"in "
                        f"{row['R_Number']} / "
                        f"{row['AI_Model']} / "
                        f"item "
                        f"{row['Item_Number']}"
                    )

            answer = row[
                "Correct_Answer"
            ]

            if answer not in {
                "A", "B", "C", "D"
            }:
                raise CommandError(
                    "Invalid correct "
                    f"answer '{answer}' "
                    f"in "
                    f"{row['R_Number']} / "
                    f"{row['AI_Model']} / "
                    f"item "
                    f"{row['Item_Number']}"
                )

            expected_skill = (
                EXPECTED_SKILLS.get(
                    row["Item_Number"]
                )
            )

            if expected_skill is None:
                raise CommandError(
                    "Invalid item number "
                    f"{row['Item_Number']}."
                )

            if (
                row["Skill"]
                != expected_skill
            ):
                raise CommandError(
                    "Skill sequence "
                    "mismatch: "
                    f"{row['R_Number']} / "
                    f"{row['AI_Model']} / "
                    f"item "
                    f"{row['Item_Number']} "
                    f"is "
                    f"'{row['Skill']}', "
                    f"expected "
                    f"'{expected_skill}'."
                )

            slot = row["RUN_SLOT"]

            if slot not in EXPECTED_RUNS:
                raise CommandError(
                    "Invalid RUN_SLOT "
                    f"{slot} in "
                    f"{row['R_Number']} / "
                    f"item "
                    f"{row['Item_Number']}."
                )

            expected_run = (
                EXPECTED_RUNS[slot]
            )

            if (
                row["AI_Model"]
                != expected_run[
                    "ai_model"
                ]
            ):
                raise CommandError(
                    "AI_Model / RUN_SLOT "
                    "mismatch in "
                    f"{row['R_Number']} / "
                    f"item "
                    f"{row['Item_Number']}: "
                    f"slot {slot} expects "
                    f"{expected_run['ai_model']!r}, "
                    f"found "
                    f"{row['AI_Model']!r}."
                )

            if (
                row["SERVICE"]
                != expected_run[
                    "service"
                ]
            ):
                raise CommandError(
                    "SERVICE / RUN_SLOT "
                    "mismatch in "
                    f"{row['R_Number']} / "
                    f"item "
                    f"{row['Item_Number']}: "
                    f"slot {slot} expects "
                    f"{expected_run['service']!r}, "
                    f"found "
                    f"{row['SERVICE']!r}."
                )

            if (
                row["CONDITION"]
                != expected_run[
                    "condition"
                ]
            ):
                raise CommandError(
                    "CONDITION mismatch "
                    f"in "
                    f"{row['R_Number']} / "
                    f"{row['SERVICE']} / "
                    f"item "
                    f"{row['Item_Number']}: "
                    f"expected "
                    f"{expected_run['condition']!r}, "
                    f"found "
                    f"{row['CONDITION']!r}."
                )

            grouped[
                (
                    row["R_Number"],
                    slot,
                )
            ].append(row)

        # Unique Round / Slot / Item key.
        duplicate_keys = [
            key
            for key, count
            in Counter(
                (
                    row["R_Number"],
                    row["RUN_SLOT"],
                    row["Item_Number"],
                )
                for row in rows
            ).items()
            if count > 1
        ]

        if duplicate_keys:
            raise CommandError(
                "Duplicate "
                "Round/RunSlot/Item "
                f"keys found: "
                f"{duplicate_keys[:20]}"
            )

        self.stdout.write(
            "\nRUN COUNTS / FROZEN CONDITIONS"
        )

        for round_id in EXPECTED_ROUNDS:
            for slot in range(1, 9):
                model_rows = grouped[
                    (round_id, slot)
                ]

                expected_run = (
                    EXPECTED_RUNS[slot]
                )

                if len(model_rows) != 10:
                    raise CommandError(
                        f"{round_id} / "
                        f"slot {slot} / "
                        f"{expected_run['service']} "
                        "must have 10 items; "
                        f"found "
                        f"{len(model_rows)}."
                    )

                item_numbers = sorted(
                    row["Item_Number"]
                    for row in model_rows
                )

                if item_numbers != (
                    list(range(1, 11))
                ):
                    raise CommandError(
                        f"{round_id} / "
                        f"slot {slot} "
                        "item numbers "
                        f"invalid: "
                        f"{item_numbers}"
                    )

                # All metadata inside a run
                # must be identical.
                for field in [
                    "AI_Model",
                    "SERVICE",
                    "CONDITION",
                    "Raw_Format",
                    "JSON_Valid",
                    "Schema_Compliant",
                ]:
                    values = {
                        row[field]
                        for row in model_rows
                    }

                    if len(values) != 1:
                        raise CommandError(
                            f"{round_id} / "
                            f"slot {slot} has "
                            f"inconsistent "
                            f"{field}: {values}"
                        )

            self.stdout.write(
                f"  {round_id}: "
                "8 frozen run slots / "
                "80 items -> PASS"
            )

        # Experiment metadata must be
        # internally consistent per round.
        for round_id in EXPECTED_ROUNDS:
            source_rows = [
                row
                for row in rows
                if row["R_Number"]
                == round_id
            ]

            for field in [
                "Domain",
                "Topic",
                "Protocol",
                "Experiment_ID",
                "CEFR",
                "Cognitive_Level",
            ]:
                values = {
                    str(
                        row[field]
                    ).strip()
                    for row
                    in source_rows
                }

                if len(values) != 1:
                    raise CommandError(
                        f"{round_id} has "
                        f"inconsistent "
                        f"{field}: {values}"
                    )

        # --------------------------------------------------
        # Provenance / recovery integrity
        # --------------------------------------------------
        recovery_counts = Counter(
            row["Recovery_Status"]
            for row in rows
        )

        if (
            dict(recovery_counts)
            != EXPECTED_RECOVERY_COUNTS
        ):
            raise CommandError(
                "Unexpected recovery "
                "provenance counts. "
                f"Expected "
                f"{EXPECTED_RECOVERY_COUNTS}, "
                f"found "
                f"{dict(recovery_counts)}"
            )

        json_run_count = 0
        non_json_run_count = 0

        for round_id in EXPECTED_ROUNDS:
            for slot in range(1, 9):
                model_rows = grouped[
                    (round_id, slot)
                ]

                raw_format = model_rows[
                    0
                ]["Raw_Format"]
                json_valid = model_rows[
                    0
                ]["JSON_Valid"]
                schema_compliant = (
                    model_rows[
                        0
                    ][
                        "Schema_Compliant"
                    ]
                )

                is_claude_interactive = (
                    slot == 4
                    and round_id
                    in {"R005", "R006"}
                )

                if is_claude_interactive:
                    if raw_format != (
                        "INTERACTIVE_QUIZ_TRANSCRIBED"
                    ):
                        raise CommandError(
                            f"{round_id} / "
                            "Claude must use "
                            "INTERACTIVE_QUIZ_TRANSCRIBED; "
                            f"found "
                            f"{raw_format!r}."
                        )

                    if (
                        json_valid
                        or schema_compliant
                    ):
                        raise CommandError(
                            f"{round_id} / "
                            "Claude must have "
                            "JSON_Valid=False and "
                            "Schema_Compliant=False."
                        )

                    for row in model_rows:
                        if (
                            row[
                                "Source_Evidence"
                            ]
                            not in (
                                None,
                                "",
                            )
                        ):
                            raise CommandError(
                                f"{round_id} / "
                                "Claude item "
                                f"{row['Item_Number']} "
                                "must preserve blank "
                                "Source_Evidence."
                            )

                    non_json_run_count += 1

                else:
                    if (
                        raw_format != "JSON"
                        or not json_valid
                        or not schema_compliant
                    ):
                        raise CommandError(
                            f"{round_id} / "
                            f"{EXPECTED_RUNS[slot]['service']} "
                            "must be strict JSON "
                            "with JSON/schema PASS. "
                            f"Found format="
                            f"{raw_format!r}, "
                            f"json_valid="
                            f"{json_valid}, "
                            f"schema="
                            f"{schema_compliant}."
                        )

                    json_run_count += 1

        self.stdout.write(
            "\nPROVENANCE / FORMAT QC"
        )
        self.stdout.write(
            "  ORIGINAL_COMBINED: "
            f"{recovery_counts['ORIGINAL_COMBINED']}"
        )
        self.stdout.write(
            "  RECOVERED_ORIGINAL_CHAT: "
            f"{recovery_counts['RECOVERED_ORIGINAL_CHAT']}"
        )
        self.stdout.write(
            "  Strict JSON runs: "
            f"{json_run_count}/40"
        )
        self.stdout.write(
            "  Interactive exact-transcription runs: "
            f"{non_json_run_count}/40"
        )

        if (
            json_run_count != 38
            or non_json_run_count != 2
        ):
            raise CommandError(
                "Frozen study expected "
                "38 strict JSON runs and "
                "2 interactive transcription runs."
            )

        # --------------------------------------------------
        # Answer-position QC
        # --------------------------------------------------
        self.stdout.write(
            "\nANSWER-POSITION QC"
        )

        pass_count = 0

        for round_id in EXPECTED_ROUNDS:
            for slot in range(1, 9):
                model_rows = grouped[
                    (round_id, slot)
                ]

                counts = Counter(
                    row["Correct_Answer"]
                    for row in model_rows
                )

                position_pass = max(
                    counts.get(
                        letter, 0
                    )
                    for letter in "ABCD"
                ) <= 4

                if position_pass:
                    pass_count += 1

                distribution = " ".join(
                    f"{letter}"
                    f"{counts.get(letter, 0)}"
                    for letter in "ABCD"
                )

                self.stdout.write(
                    f"  {round_id} / "
                    f"slot {slot} / "
                    f"{EXPECTED_RUNS[slot]['service']}: "
                    f"{distribution} -> "
                    f"{'PASS' if position_pass else 'FAIL'}"
                )

        self.stdout.write(
            "\nAnswer-position "
            f"compliant runs: "
            f"{pass_count}/40"
        )

        if pass_count != 15:
            raise CommandError(
                "Frozen historical QC "
                "requires exactly 15/40 "
                "answer-position-compliant runs; "
                f"found {pass_count}/40."
            )

    @transaction.atomic
    def _import(self, rows):
        if (
            ResearchExperiment.objects.exists()
            or ResearchRun.objects.exists()
            or ResearchQuestion.objects.exists()
        ):
            raise CommandError(
                "Research tables are not empty. "
                "Import aborted to protect "
                "frozen research data."
            )

        grouped = defaultdict(list)

        for row in rows:
            grouped[
                (
                    row["R_Number"],
                    row["RUN_SLOT"],
                )
            ].append(row)

        for round_id in EXPECTED_ROUNDS:
            round_rows = [
                row
                for row in rows
                if row["R_Number"]
                == round_id
            ]

            first = round_rows[0]

            experiment = (
                ResearchExperiment.objects.create(
                    experiment_id=first[
                        "Experiment_ID"
                    ],
                    source_id=round_id,
                    domain=first[
                        "Domain"
                    ],
                    topic=first[
                        "Topic"
                    ],
                    protocol=first[
                        "Protocol"
                    ],
                    target_cefr=first[
                        "CEFR"
                    ],
                    target_cognitive_level=first[
                        "Cognitive_Level"
                    ],
                    is_confirmatory=True,
                    is_frozen=True,
                )
            )

            for slot in range(1, 9):
                model_rows = sorted(
                    grouped[
                        (round_id, slot)
                    ],
                    key=lambda row: (
                        row[
                            "Item_Number"
                        ]
                    ),
                )

                first_run_row = (
                    model_rows[0]
                )

                counts = Counter(
                    row[
                        "Correct_Answer"
                    ]
                    for row
                    in model_rows
                )

                answer_position_pass = (
                    max(
                        counts.get(
                            letter, 0
                        )
                        for letter
                        in "ABCD"
                    )
                    <= 4
                )

                run_meta = (
                    EXPECTED_RUNS[slot]
                )

                run = (
                    ResearchRun.objects.create(
                        experiment=experiment,
                        run_id=(
                            f"{round_id}-V4-"
                            f"RUN-{slot:03d}"
                        ),
                        provider=first_run_row[
                            "SERVICE"
                        ],
                        displayed_model=run_meta[
                            "displayed_model"
                        ],
                        condition=first_run_row[
                            "CONDITION"
                        ],
                        raw_output_format=first_run_row[
                            "Raw_Format"
                        ],
                        json_valid=first_run_row[
                            "JSON_Valid"
                        ],
                        schema_compliant=first_run_row[
                            "Schema_Compliant"
                        ],
                        item_count_compliant=True,
                        skill_sequence_compliant=True,
                        answer_position_compliant=(
                            answer_position_pass
                        ),
                    )
                )

                for row in model_rows:
                    ResearchQuestion.objects.create(
                        run=run,
                        # IMPORTANT:
                        # consume the frozen,
                        # pre-existing blind ID.
                        # Never regenerate it.
                        blind_id=row[
                            "QUESTION_ID"
                        ],
                        item_number=row[
                            "Item_Number"
                        ],
                        skill=row[
                            "Skill"
                        ],
                        stem=row[
                            "Stem"
                        ],
                        option_a=row[
                            "Option_A"
                        ],
                        option_b=row[
                            "Option_B"
                        ],
                        option_c=row[
                            "Option_C"
                        ],
                        option_d=row[
                            "Option_D"
                        ],
                        correct_answer=row[
                            "Correct_Answer"
                        ],
                        explanation=(
                            row[
                                "Explanation"
                            ]
                            or ""
                        ),
                        source_evidence=(
                            row[
                                "Source_Evidence"
                            ]
                            or ""
                        ),
                        raw_preserved=True,
                        is_validated=False,
                        promoted_to_operational_bank=False,
                    )

        # --------------------------------------------------
        # Post-import hard checks
        # --------------------------------------------------
        experiment_count = (
            ResearchExperiment.objects.count()
        )
        run_count = (
            ResearchRun.objects.count()
        )
        question_count = (
            ResearchQuestion.objects.count()
        )

        if experiment_count != 5:
            raise CommandError(
                "Post-import validation "
                "failed: expected 5 "
                f"experiments, found "
                f"{experiment_count}."
            )

        if run_count != 40:
            raise CommandError(
                "Post-import validation "
                "failed: expected 40 "
                f"runs, found "
                f"{run_count}."
            )

        if question_count != 400:
            raise CommandError(
                "Post-import validation "
                "failed: expected 400 "
                f"questions, found "
                f"{question_count}."
            )

        imported_ids = set(
            ResearchQuestion.objects.values_list(
                "blind_id",
                flat=True,
            )
        )

        if imported_ids != (
            EXPECTED_QUESTION_IDS
        ):
            raise CommandError(
                "Post-import validation "
                "failed: database "
                "Blind-ID set does not "
                "exactly match "
                "AEV4-0001..AEV4-0400."
            )

        # Verify each frozen question maps
        # back to its intended round/slot/item.
        for row in rows:
            question = (
                ResearchQuestion.objects
                .select_related(
                    "run",
                    "run__experiment",
                )
                .get(
                    blind_id=row[
                        "QUESTION_ID"
                    ]
                )
            )

            expected_run_id = (
                f"{row['R_Number']}"
                f"-V4-RUN-"
                f"{row['RUN_SLOT']:03d}"
            )

            if (
                question.run.run_id
                != expected_run_id
                or question.item_number
                != row["Item_Number"]
                or question.run.experiment.source_id
                != row["R_Number"]
            ):
                raise CommandError(
                    "Post-import Blind-ID "
                    "mapping mismatch for "
                    f"{row['QUESTION_ID']}."
                )

        self.stdout.write(
            self.style.SUCCESS(
                "\nIMPORT COMPLETE - "
                "FROZEN V4 DATASET INSTALLED"
            )
        )
        self.stdout.write(
            f"Experiments: "
            f"{experiment_count}"
        )
        self.stdout.write(
            f"Runs: {run_count}"
        )
        self.stdout.write(
            f"Questions: "
            f"{question_count}"
        )
        self.stdout.write(
            "Blind IDs: "
            "400/400 preserved exactly"
        )
        self.stdout.write(
            "Operational question bank: "
            "not touched by this importer"
        )
